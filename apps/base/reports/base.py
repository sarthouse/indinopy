import csv
import io
import logging
from abc import ABC, abstractmethod
from typing import Any, Dict, List, Optional
from django.http import HttpResponse
from django.template.loader import render_to_string

logger = logging.getLogger(__name__)


class BaseReport(ABC):
    """
    Clase base abstracta universal para cualquier generador de reportes en Indinopy.
    Desacopla la extracción y procesamiento de datos del formato de salida final.
    """

    nombre_archivo: str = "reporte"
    content_type: str = "text/plain"
    extension: str = "txt"

    def __init__(self, objeto_origen: Any = None, request: Optional[Any] = None, **kwargs):
        self.objeto = objeto_origen
        self.request = request
        self.params = kwargs

    @abstractmethod
    def render(self) -> bytes:
        """Genera el contenido binario o de texto del reporte en memoria."""
        pass

    def get_filename(self) -> str:
        """Devuelve el nombre completo del archivo con su extensión correspondiente."""
        base = self.nombre_archivo.rstrip(f".{self.extension}")
        return f"{base}.{self.extension}"

    def to_http_response(self, inline: bool = False) -> HttpResponse:
        """
        Empaqueta el contenido generado en un HttpResponse listo para el navegador.
        - inline=True: Abre en visor del navegador si está soportado (ej. PDF).
        - inline=False: Fuerza descarga como archivo adjunto (Attachment).
        """
        contenido = self.render()
        response = HttpResponse(contenido, content_type=self.content_type)
        disposition = "inline" if inline else "attachment"
        filename = self.get_filename()
        response["Content-Disposition"] = f'{disposition}; filename="{filename}"'
        return response


class BasePDFReport(BaseReport):
    """
    Generador de reportes imprimibles en PDF (ej. Facturas, Remitos, Hoja de Ruta e-OP).
    Utiliza el motor de templates HTML/CSS de Django y lo convierte a PDF (WeasyPrint o fallback HTML).
    """

    template_name: str
    content_type: str = "application/pdf"
    extension: str = "pdf"

    @abstractmethod
    def get_context_data(self) -> Dict[str, Any]:
        """Recolecta los datos necesarios para renderizar el template."""
        pass

    def render_html(self) -> str:
        """Renderiza el template Django con el contexto estructurado."""
        context = self.get_context_data()
        context["report_obj"] = self.objeto
        context["params"] = self.params
        return render_to_string(self.template_name, context, request=self.request)

    def render(self) -> bytes:
        html_content = self.render_html()

        # Intentar renderizar con WeasyPrint
        try:
            from weasyprint import HTML
            base_url = self.request.build_absolute_uri('/') if self.request else None
            return HTML(string=html_content, base_url=base_url).write_pdf()
        except ImportError:
            logger.warning(
                "WeasyPrint no está instalado. Retornando HTML codificado en UTF-8 como fallback de impresión."
            )
            return html_content.encode("utf-8")


class BaseTabularReport(BaseReport):
    """
    Generador de reportes de datos tabulares (Libro IVA, Listados de Stock, Balances).
    Soporta exportación a CSV nativo y Excel (.xlsx) con columnas configurables.
    """

    content_type: str = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension: str = "xlsx"

    @abstractmethod
    def get_headers(self) -> List[str]:
        """Lista de títulos de columnas (ej. ['Fecha', 'Comprobante', 'CUIT', 'Neto Gravado', 'IVA'])."""
        pass

    @abstractmethod
    def get_rows(self) -> List[List[Any]]:
        """Matriz de datos (lista de filas donde cada fila es una lista de valores ordenados)."""
        pass

    def render_csv(self, delimiter: str = ";") -> bytes:
        """Genera el contenido en formato CSV compatible con Excel latinoamericano (punto y coma)."""
        output = io.StringIO()
        # Escribir BOM UTF-8 para que Excel en Windows abra tildes y caracteres correctamente
        output.write("\ufeff")
        writer = csv.writer(output, delimiter=delimiter)

        writer.writerow(self.get_headers())
        for row in self.get_rows():
            # Limpiar formatos de datos para CSV
            cleaned_row = [str(val) if val is not None else "" for val in row]
            writer.writerow(cleaned_row)

        return output.getvalue().encode("utf-8")

    def render_excel(self) -> bytes:
        """
        Genera el libro Excel con estilos básicos de cabecera y ajuste de columnas.
        Utiliza openpyxl si está disponible, o realiza fallback automático a CSV.
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte"

            # 1. Cabeceras con estilo corporativo
            headers = self.get_headers()
            ws.append(headers)

            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1")
            )

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # 2. Filas de datos
            rows = self.get_rows()
            for row in rows:
                ws.append(row)

            # 3. Auto-ajuste de ancho de columnas
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            output = io.BytesIO()
            wb.save(output)
            return output.getvalue()

        except ImportError:
            logger.info("openpyxl no está instalado. Generando fallback a CSV compatible con Excel.")
            self.content_type = "text/csv; charset=utf-8"
            self.extension = "csv"
            return self.render_csv()

    def render(self) -> bytes:
        if self.extension == "csv":
            return self.render_csv()
        return self.render_excel()
